import datetime
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from utils.config import CLASS_LABELS
from utils.logger import logger

NAVY = RGBColor(0x1F, 0x38, 0x64)
MUTED = RGBColor(0x6B, 0x75, 0x87)


def set_cell_margins(cell, **kwargs):
    """Sets cell margins for a docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m in ['top', 'start', 'bottom', 'end']:
        if m in kwargs:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(kwargs[m]))
            node.set(qn('w:type'), 'dxa')
            tc_mar.append(node)
    tc_pr.append(tc_mar)


def _table(doc, rows, header=None):
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.style = 'Table Grid'
    if header:
        cells = table.add_row().cells
        for cell, text in zip(cells, header):
            cell.text = text
            cell.paragraphs[0].runs[0].font.bold = True
            set_cell_margins(cell, start=100)
    for row in rows:
        cells = table.add_row().cells
        for i, (cell, text) in enumerate(zip(cells, row)):
            cell.text = str(text)
            if i == 0 and not header:
                cell.paragraphs[0].runs[0].font.bold = True
            set_cell_margins(cell, start=100)
    return table


def _picture(doc, buffer, caption, width=3.2):
    buffer.seek(0)
    doc.add_picture(buffer, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = MUTED


def generate_docx_report(image_buffer, prediction, probabilities, patient_id, scan_id,
                         model_version, gradcam_buffer=None, similar_cases=None, study_id=None, slice_index=None):
    """
    Generates a template-based brain tumor MRI classification report (DOCX).

    prediction:     class key, e.g. "glioma"
    probabilities:  {"glioma": 0.93, "meningioma": 0.04, "notumor": 0.01, "pituitary": 0.02}
    similar_cases:  optional list of {"case_id", "label", "similarity"} (similarity 0-100)
    study_id:       set when the scan is one slice of a multi-slice study (slice_index is its position)
    """
    logger.info(f"Generating report for scan {scan_id}: {prediction}")
    doc = Document()

    # 1. Header
    header_para = doc.sections[0].header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_para.add_run("NEUROLENS | BRAIN TUMOR MRI CLASSIFICATION SUPPORT")
    run.font.size = Pt(9)
    run.font.color.rgb = NAVY

    # 2. Title & metadata
    title = doc.add_heading('Brain MRI Classification Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"Generated {datetime.datetime.now().strftime('%d %B %Y, %H:%M')}")
    run.font.size = Pt(10)
    run.font.italic = True

    # 3. Summary
    doc.add_heading('1. Summary', level=1)
    confidence = probabilities.get(prediction, 0.0)
    summary = _table(doc, [
        ("Patient", patient_id),
        ("Scan", scan_id),
        ("Scan type", f"Slice {slice_index} of multi-slice study {study_id}" if study_id else "Single slice"),
        ("Predicted class", CLASS_LABELS.get(prediction, prediction)),
        ("Confidence", f"{confidence * 100:.1f}%"),
        ("Model", model_version),
    ])
    summary.rows[3].cells[1].paragraphs[0].runs[0].font.bold = True
    summary.rows[3].cells[1].paragraphs[0].runs[0].font.color.rgb = NAVY
    if confidence < 0.7:
        p = doc.add_paragraph()
        r = p.add_run("Low confidence (< 70%) — this result should be reviewed with particular care.")
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xC9, 0x81, 0x1A)

    # 4. Class probabilities
    doc.add_heading('2. Class probabilities', level=1)
    ordered = sorted(probabilities.items(), key=lambda kv: kv[1], reverse=True)
    _table(doc, [(CLASS_LABELS.get(c, c), f"{p * 100:.1f}%") for c, p in ordered], header=("Class", "Probability"))

    # 5. Images
    doc.add_heading('3. Scan and explanation', level=1)
    if image_buffer:
        _picture(doc, image_buffer, "Uploaded MRI slice")
    if gradcam_buffer:
        doc.add_paragraph(
            "The heatmap highlights the regions (warmer colours) that contributed most to the predicted class. "
            "It shows where the model looked; it is not a segmentation of the tumor."
        )
        _picture(doc, gradcam_buffer, "Explanation heatmap")

    # 6. Similar confirmed cases
    if similar_cases:
        doc.add_heading('4. Most similar confirmed cases', level=1)
        _table(doc, [(c["case_id"], CLASS_LABELS.get(c["label"], c["label"]), f"{c['similarity']:.1f}%") for c in similar_cases],
               header=("Case", "Confirmed class", "Similarity"))
        agree = sum(1 for c in similar_cases if c["label"] == prediction)
        doc.add_paragraph(f"{agree} of the {len(similar_cases)} most similar confirmed cases share the predicted class.")

    # 7. Disclaimer
    return _finish(doc)


def _finish(doc):
    """Adds the research-prototype disclaimer and returns the document as an in-memory DOCX file."""
    doc.add_paragraph()
    disclaimer = doc.add_paragraph()
    disclaimer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = disclaimer.add_run("RESEARCH PROTOTYPE — NOT A DIAGNOSTIC DEVICE")
    run.bold = True
    run.font.size = Pt(11)
    p = doc.add_paragraph(
        "This report was generated automatically by a research prototype to support, not replace, clinical judgement. "
        "All findings must be reviewed and verified by a qualified radiologist."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.size = Pt(9)
        r.font.italic = True

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def generate_study_report(study, slices):
    """
    Report for a multi-slice study (DOCX).

    study:  combined result — id, patient_id, prediction, confidence, probabilities, slice_agreement, model_version
    slices: list of (scan_record, image_buffer_or_None, heatmap_buffer_or_None), in slice order
    """
    logger.info(f"Generating report for study {study['id']}: {study['prediction']}")
    doc = Document()
    header_para = doc.sections[0].header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_para.add_run("NEUROLENS | BRAIN TUMOR MRI CLASSIFICATION SUPPORT")
    run.font.size = Pt(9)
    run.font.color.rgb = NAVY

    title = doc.add_heading('Multi-slice MRI Classification Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"Generated {datetime.datetime.now().strftime('%d %B %Y, %H:%M')}")
    run.font.size = Pt(10)
    run.font.italic = True

    doc.add_heading('1. Combined result', level=1)
    summary = _table(doc, [
        ("Patient", study["patient_id"]),
        ("Study", study["id"]),
        ("Scan type", f"Multi-slice study · {len(slices)} slices"),
        ("Combined prediction", CLASS_LABELS.get(study["prediction"], study["prediction"])),
        ("Average confidence", f"{study['confidence'] * 100:.1f}%"),
        ("Slices that agree", study["slice_agreement"]),
        ("Model", study.get("model_version", "")),
    ])
    summary.rows[3].cells[1].paragraphs[0].runs[0].font.bold = True
    summary.rows[3].cells[1].paragraphs[0].runs[0].font.color.rgb = NAVY
    if study.get("status") == "Flagged":
        p = doc.add_paragraph()
        r = p.add_run("Flagged for careful review: the slices disagree or the average confidence is below 70%.")
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xC9, 0x81, 0x1A)
    doc.add_paragraph(
        "Each slice is classified separately by the same 2D model, and the class probabilities are averaged. "
        "The model does not analyse the scan in 3D.")

    doc.add_heading('2. Average class probabilities', level=1)
    ordered = sorted(study["probabilities"].items(), key=lambda kv: kv[1], reverse=True)
    _table(doc, [(CLASS_LABELS.get(c, c), f"{p * 100:.1f}%") for c, p in ordered], header=("Class", "Average probability"))

    doc.add_heading('3. Per-slice results', level=1)
    _table(doc, [(str(scan.get("slice_index", i + 1)), scan["id"], CLASS_LABELS.get(scan["prediction"], scan["prediction"]),
                  f"{scan['confidence'] * 100:.1f}%") for i, (scan, _, _) in enumerate(slices)],
           header=("Slice", "Scan", "Prediction", "Confidence"))

    doc.add_heading('4. Slices and explanations', level=1)
    doc.add_paragraph("Left: uploaded slice. Right: heatmap of the regions that contributed most to that slice's prediction.")
    grid = doc.add_table(rows=0, cols=2)
    for i, (scan, image, heatmap) in enumerate(slices):
        caption = (f"Slice {scan.get('slice_index', i + 1)} · {scan['id']} · "
                   f"{CLASS_LABELS.get(scan['prediction'], scan['prediction'])} ({scan['confidence'] * 100:.1f}%)")
        cells = grid.add_row().cells
        for cell, buf in zip(cells, (image, heatmap)):
            if buf:
                buf.seek(0)
                cell.paragraphs[0].add_run().add_picture(buf, width=Inches(2.4))
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_cells = grid.add_row().cells
        label_cells[0].merge(label_cells[1])
        label_cells[0].text = caption
        label_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_cells[0].paragraphs[0].runs[0].font.size = Pt(9)
        label_cells[0].paragraphs[0].runs[0].font.color.rgb = MUTED

    return _finish(doc)
